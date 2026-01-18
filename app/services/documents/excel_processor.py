"""
معالج Excel
Excel Document Processor (XLSX/XLS)
"""

from pathlib import Path
from typing import Dict, Any, List
import pandas as pd
import openpyxl

from .base_processor import BaseDocumentProcessor, ProcessedDocument


class ExcelProcessor(BaseDocumentProcessor):
    """معالج مستندات Excel"""

    supported_extensions = [".xlsx", ".xls", ".csv"]
    processor_name = "excel"

    async def process(self, file_path: Path) -> ProcessedDocument:
        """معالجة Excel واستخراج البيانات"""

        extension = file_path.suffix.lower()

        if extension == ".csv":
            df_dict = {"Sheet1": pd.read_csv(str(file_path))}
        else:
            df_dict = pd.read_excel(str(file_path), sheet_name=None)

        tables = []
        text_parts = []

        for sheet_name, df in df_dict.items():
            # Convert to text representation
            sheet_text = f"--- {sheet_name} ---\n"
            sheet_text += df.to_string(index=False)
            text_parts.append(sheet_text)

            # Store table data
            tables.append({
                "sheet_name": sheet_name,
                "rows": len(df),
                "columns": len(df.columns),
                "column_names": df.columns.tolist(),
                "data_preview": df.head(10).to_dict(orient="records")
            })

        combined_text = "\n\n".join(text_parts)
        metadata = await self.extract_metadata(file_path)

        return ProcessedDocument(
            text=combined_text,
            pages=text_parts,
            page_count=len(df_dict),
            word_count=self.count_words(combined_text),
            language=self.detect_language(combined_text),
            metadata=metadata,
            tables=tables,
            images=[]
        )

    async def extract_metadata(self, file_path: Path) -> Dict[str, Any]:
        """استخراج البيانات الوصفية"""

        if file_path.suffix.lower() == ".csv":
            return {"type": "csv", "sheets": 1}

        wb = openpyxl.load_workbook(str(file_path), read_only=True)

        return {
            "sheets": wb.sheetnames,
            "sheet_count": len(wb.sheetnames),
            "created": wb.properties.created.isoformat() if wb.properties.created else None,
            "modified": wb.properties.modified.isoformat() if wb.properties.modified else None,
            "creator": wb.properties.creator,
            "title": wb.properties.title
        }
